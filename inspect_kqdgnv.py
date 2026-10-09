import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('BẢNG THIẾT LẬP TIÊU CHÍ ĐÁNH GIÁ NHÂN VIÊN (Responses) (1).xlsx', data_only=True)
ws = wb['KQDGNV']

headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
for i, h in enumerate(headers, 1):
    val2 = ws.cell(row=2, column=i).value
    val3 = ws.cell(row=3, column=i).value
    val4 = ws.cell(row=4, column=i).value
    print(f"Col {i:2d} | Header: {str(h):40s} | Row2: {str(val2):20s} | Row3: {str(val3):20s}")
