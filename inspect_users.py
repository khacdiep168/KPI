import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('BẢNG THIẾT LẬP TIÊU CHÍ ĐÁNH GIÁ NHÂN VIÊN (Responses) (1).xlsx', data_only=True)

print("=== Users Sheet ===")
ws = wb['Users']
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if any(vals):
        print(f"Row {r:2d}: {vals}")

print("\n=== DSNV Sheet ===")
ws = wb['DSNV']
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if any(vals):
        print(f"Row {r:2d}: {vals}")
