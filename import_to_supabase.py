import openpyxl
import json
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BẢNG THIẾT LẬP TIÊU CHÍ ĐÁNH GIÁ NHÂN VIÊN (Responses) (1).xlsx', data_only=True)

# 1. Inspect Users
ws_users = wb['Users']
users = []
for r in range(2, ws_users.max_row + 1):
    manv = ws_users.cell(row=r, column=1).value
    name = ws_users.cell(row=r, column=2).value
    dept = ws_users.cell(row=r, column=3).value
    role = ws_users.cell(row=r, column=4).value
    pwd = ws_users.cell(row=r, column=5).value
    if manv:
        users.append({
            'manv': str(manv).strip(),
            'name': str(name).strip() if name else '',
            'dept': str(dept).strip() if dept else '',
            'role': str(role).strip() if role else 'User',
            'password': str(pwd).strip() if pwd else '123'
        })

print(f"Users found in 'Users' sheet: {len(users)}")

# 2. Inspect KQDGNV
ws_kq = wb['KQDGNV']
print(f"KQDGNV max_row: {ws_kq.max_row}")
headers = [ws_kq.cell(row=1, column=c).value for c in range(1, ws_kq.max_column + 1)]
print("Headers in KQDGNV:", headers)

unique_months = set()
valid_evals = 0
for r in range(2, ws_kq.max_row + 1):
    manv = ws_kq.cell(row=r, column=2).value
    month = ws_kq.cell(row=r, column=25).value
    if manv:
        valid_evals += 1
        if month:
            unique_months.add(str(month).strip())

print(f"Valid evaluations in KQDGNV: {valid_evals}")
print(f"Unique months/periods found: {unique_months}")
