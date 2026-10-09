import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BẢNG THIẾT LẬP TIÊU CHÍ ĐÁNH GIÁ NHÂN VIÊN (Responses) (1).xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)

for s in wb.sheetnames:
    ws = wb[s]
    headers = [cell.value for cell in ws[1]]
    print(f"\n=== Sheet: {s} (rows: {ws.max_row}, cols: {ws.max_column}) ===")
    print("Headers:", [h for h in headers if h is not None])
    print("Sample row 2:", [cell.value for cell in ws[2]][:8])
    if ws.max_row >= 3:
        print("Sample row 3:", [cell.value for cell in ws[3]][:8])
