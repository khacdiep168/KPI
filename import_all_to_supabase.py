import openpyxl
import urllib.request
import json
import re
import sys
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

SUPABASE_URL = "https://jaewyxazjvqtvqifdmlw.supabase.co"
ANON_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImphZXd5eGF6anZxdHZxaWZkbWx3Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTE0Mzg0MTgsImV4cCI6MjEwNzAxNDQxOH0.vXxjYEdEvK7HCefzPuKXgcguB4KPRjGkJhHcPF9uW00"

def post_batch(table, items):
    if not items:
        return
    url = f"{SUPABASE_URL}/rest/v1/{table}?on_conflict=manv" if table == 'users' else f"{SUPABASE_URL}/rest/v1/{table}"
    req = urllib.request.Request(
        url,
        data=json.dumps(items).encode('utf-8'),
        headers={
            "apikey": ANON_KEY,
            "Authorization": f"Bearer {ANON_KEY}",
            "Content-Type": "application/json",
            "Prefer": "resolution=merge-duplicates,return=minimal"
        },
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        if resp.status not in (200, 201, 204):
            print(f"Failed to post to {table}, status: {resp.status}")

wb = openpyxl.load_workbook('BẢNG THIẾT LẬP TIÊU CHÍ ĐÁNH GIÁ NHÂN VIÊN (Responses) (1).xlsx', data_only=True)

# ==========================================
# 1. PARSE USERS
# ==========================================
print("\n--- 1. PARSING USERS ---")
users_dict = {}

# From Users sheet
ws_users = wb['Users']
for r in range(2, ws_users.max_row + 1):
    raw_manv = ws_users.cell(row=r, column=1).value
    if not raw_manv:
        continue
    manv = str(raw_manv).replace('\xa0', '').strip().upper()
    name = str(ws_users.cell(row=r, column=2).value or '').strip()
    dept = str(ws_users.cell(row=r, column=3).value or '').strip()
    raw_role = str(ws_users.cell(row=r, column=4).value or '').strip()
    pwd = str(ws_users.cell(row=r, column=5).value or '123').strip()
    
    role = 'User'
    if raw_role.lower() in ['admin']:
        role = 'Admin'
    elif raw_role.lower() in ['quanly', 'manager', 'quản lý']:
        role = 'Manager'
        
    title = 'NHÂN VIÊN'
    if role == 'Admin' and 'GIÁM ĐỐC' in name.upper():
        title = 'GIÁM ĐỐC'
    elif role in ['Admin', 'Manager']:
        title = 'QUẢN LÝ'
        
    users_dict[manv] = {
        'manv': manv,
        'name': name,
        'dept': dept,
        'title': title,
        'password': pwd,
        'role': role
    }

# From DSNV sheet (fill in any missing)
ws_dsnv = wb['DSNV']
for r in range(2, ws_dsnv.max_row + 1):
    raw_manv = ws_dsnv.cell(row=r, column=1).value
    if not raw_manv:
        continue
    manv = str(raw_manv).replace('\xa0', '').strip().upper()
    name = str(ws_dsnv.cell(row=r, column=2).value or '').strip()
    dept = str(ws_dsnv.cell(row=r, column=3).value or '').strip()
    chuc_vu = str(ws_dsnv.cell(row=r, column=4).value or 'NHÂN VIÊN').strip()
    pwd = str(ws_dsnv.cell(row=r, column=5).value or '123').strip()
    raw_roll = str(ws_dsnv.cell(row=r, column=6).value or '').strip()
    
    role = 'User'
    if raw_roll.lower() in ['admin']:
        role = 'Admin'
    elif raw_roll.lower() in ['quanly', 'manager', 'quản lý'] or 'QUẢN LÝ' in chuc_vu.upper():
        role = 'Manager'

    if manv not in users_dict:
        users_dict[manv] = {
            'manv': manv,
            'name': name,
            'dept': dept,
            'title': chuc_vu,
            'password': pwd,
            'role': role
        }

# Ensure Key Management Accounts are properly set
key_mgmt = {
    'BGD01': {'name': 'BAN GIÁM ĐỐC', 'dept': 'BAN GIÁM ĐỐC', 'title': 'GIÁM ĐỐC', 'password': '123', 'role': 'Admin'},
    'HC01': {'name': 'QUẢN LÝ HÀNH CHÍNH', 'dept': 'HÀNH CHÍNH', 'title': 'QUẢN LÝ', 'password': '123', 'role': 'Manager'},
    'DTK01': {'name': 'QUẢN LÝ ĐIỂM TẬP KẾT', 'dept': 'ĐIỂM TẬP KẾT', 'title': 'QUẢN LÝ', 'password': '123', 'role': 'Manager'},
    'KD01': {'name': 'QUẢN LÝ KINH DOANH', 'dept': 'PHÒNG KINH DOANH', 'title': 'QUẢN LÝ', 'password': '123', 'role': 'Manager'},
    'KT01': {'name': 'QUẢN LÝ KẾ TOÁN', 'dept': 'PHÒNG KẾ TOÁN', 'title': 'QUẢN LÝ', 'password': '123', 'role': 'Manager'},
}
for k, v in key_mgmt.items():
    if k not in users_dict:
        users_dict[k] = {'manv': k, **v}
    else:
        users_dict[k]['role'] = v['role']
        users_dict[k]['dept'] = v['dept']

user_list = list(users_dict.values())
print(f"Total clean users to import: {len(user_list)}")
post_batch('users', user_list)
print("-> Successfully imported users to Supabase!")

# Helper map: Name to user info
name_to_user = {}
for u in user_list:
    name_norm = re.sub(r'\s+', ' ', u['name']).strip().upper()
    name_to_user[name_norm] = u

# ==========================================
# 2. PARSE EVALUATIONS & PERIODS
# ==========================================
print("\n--- 2. PARSING KQDGNV EVALUATIONS ---")
ws_kq = wb['KQDGNV']

periods_set = set()
evaluations_list = []

def clean_num(val, default=0):
    if val is None or val == '':
        return default
    try:
        if isinstance(val, (int, float)):
            return float(val)
        cleaned = re.sub(r'[^\d.]', '', str(val))
        return float(cleaned) if cleaned else default
    except:
        return default

def parse_date(val):
    if not val:
        return datetime.now().isoformat()
    if isinstance(val, datetime):
        return val.isoformat()
    # Try parsing string format: D/M/Y or Y-M-D
    val_str = str(val).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%m/%d/%Y %H:%M:%S', '%d/%m/%Y', '%m/%d/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try:
            return datetime.strptime(val_str, fmt).isoformat()
        except:
            pass
    return datetime.now().isoformat()

def extract_period(month_val, date_val):
    if month_val:
        m_str = str(month_val).strip()
        # Pattern MM-YYYY
        match = re.search(r'(\d{1,2})[-/](\d{4})', m_str)
        if match:
            mm = int(match.group(1))
            yyyy = int(match.group(2))
            return f"{mm:02d}-{yyyy}"
    # Extract from date_val
    if isinstance(date_val, datetime):
        return f"{date_val.month:02d}-{date_val.year}"
    if date_val:
        d_str = str(date_val).strip()
        match = re.search(r'(\d{1,2})[/](\d{1,2})[/](\d{4})', d_str)
        if match:
            # Usually M/D/YYYY in Google Forms
            mm = int(match.group(1))
            yyyy = int(match.group(3))
            return f"{mm:02d}-{yyyy}"
    return "10-2026"

for r in range(2, ws_kq.max_row + 1):
    raw_evaluator_manv = ws_kq.cell(row=r, column=2).value
    if not raw_evaluator_manv:
        continue
    evaluator_manv = str(raw_evaluator_manv).replace('\xa0', '').strip().upper()
    evaluator_name = str(ws_kq.cell(row=r, column=3).value or '').strip()
    evaluator_dept = str(ws_kq.cell(row=r, column=4).value or '').strip()
    target_name = str(ws_kq.cell(row=r, column=5).value or '').strip()
    
    if not target_name:
        target_name = evaluator_name
        
    date_val = ws_kq.cell(row=r, column=1).value
    created_at = parse_date(date_val)
    
    month_val = ws_kq.cell(row=r, column=25).value
    period = extract_period(month_val, date_val)
    periods_set.add(period)
    
    # Resolve target_manv & target_dept
    target_norm = re.sub(r'\s+', ' ', target_name).strip().upper()
    target_manv = evaluator_manv
    target_dept = evaluator_dept
    
    if target_norm == re.sub(r'\s+', ' ', evaluator_name).strip().upper():
        target_manv = evaluator_manv
        target_dept = evaluator_dept
    elif target_norm in name_to_user:
        target_manv = name_to_user[target_norm]['manv']
        target_dept = name_to_user[target_norm]['dept']
    else:
        # Fuzzy match
        for uname, u in name_to_user.items():
            if target_norm in uname or uname in target_norm:
                target_manv = u['manv']
                target_dept = u['dept']
                break

    q1 = clean_num(ws_kq.cell(row=r, column=6).value)
    q2 = clean_num(ws_kq.cell(row=r, column=7).value)
    q3 = clean_num(ws_kq.cell(row=r, column=8).value)
    q4 = clean_num(ws_kq.cell(row=r, column=9).value)
    q5 = clean_num(ws_kq.cell(row=r, column=10).value)
    q6 = clean_num(ws_kq.cell(row=r, column=11).value)
    q7 = clean_num(ws_kq.cell(row=r, column=12).value)
    q8 = clean_num(ws_kq.cell(row=r, column=13).value)
    q9 = clean_num(ws_kq.cell(row=r, column=14).value)
    q10 = clean_num(ws_kq.cell(row=r, column=15).value)
    
    self_feedback = str(ws_kq.cell(row=r, column=16).value or '').strip()
    self_total = clean_num(ws_kq.cell(row=r, column=17).value)
    if self_total == 0:
        self_total = q1 + q2 + q3 + q4 + q5 + q6 + q7 + q8 + q9 + q10
        
    self_cls = str(ws_kq.cell(row=r, column=18).value or '').strip()
    if not self_cls:
        if self_total >= 80: self_cls = "Giỏi"
        elif self_total >= 65: self_cls = "Khá"
        elif self_total >= 50: self_cls = "Trung Bình"
        else: self_cls = "Yếu"
        
    mgr_score = ws_kq.cell(row=r, column=19).value
    mgr_score = clean_num(mgr_score, None) if mgr_score is not None else None
    mgr_cls = str(ws_kq.cell(row=r, column=20).value or '').strip() or None
    mgr_feedback = str(ws_kq.cell(row=r, column=21).value or '').strip() or None
    
    council_score = ws_kq.cell(row=r, column=22).value
    council_score = clean_num(council_score, None) if council_score is not None else None
    council_cls = str(ws_kq.cell(row=r, column=23).value or '').strip() or None
    council_feedback = str(ws_kq.cell(row=r, column=24).value or '').strip() or None
    
    status = "SUBMITTED"
    if council_score is not None:
        status = "APPROVED"
    elif mgr_score is not None:
        status = "MANAGER_REVIEWED"
        
    evaluations_list.append({
        'id': f"eval-excel-{r}",
        'period': period,
        'evaluator_manv': evaluator_manv,
        'evaluator_name': evaluator_name,
        'evaluator_dept': evaluator_dept,
        'target_manv': target_manv,
        'target_name': target_name,
        'target_dept': target_dept,
        'q1': q1,
        'q2': q2,
        'q3': q3,
        'q4': q4,
        'q5': q5,
        'q6': q6,
        'q7': q7,
        'q8': q8,
        'q9': q9,
        'q10': q10,
        'self_total_score': self_total,
        'self_classification': self_cls,
        'self_feedback': self_feedback,
        'manager_score': mgr_score,
        'manager_classification': mgr_cls,
        'manager_feedback': mgr_feedback,
        'council_score': council_score,
        'council_classification': council_cls,
        'council_feedback': council_feedback,
        'status': status,
        'created_at': created_at,
        'updated_at': created_at
    })

print(f"Total evaluations parsed: {len(evaluations_list)}")
print(f"Unique periods detected: {sorted(list(periods_set))}")

# ==========================================
# 3. IMPORT PERIODS
# ==========================================
print("\n--- 3. IMPORTING KPI PERIODS ---")
periods_list = []
for p in sorted(list(periods_set)):
    is_open = (p == "10-2026")
    periods_list.append({
        'id': f"p-{p}",
        'month_year': p,
        'name': f"Kỳ đánh giá KPI Tháng {p}",
        'is_open': is_open
    })
# Ensure 10-2026 exists
if not any(p['month_year'] == '10-2026' for p in periods_list):
    periods_list.append({
        'id': 'p-10-2026',
        'month_year': '10-2026',
        'name': 'Kỳ đánh giá KPI Tháng 10/2026',
        'is_open': True
    })

post_batch('kpi_periods', periods_list)
print("-> Successfully imported periods to Supabase!")

# ==========================================
# 4. BATCH IMPORT EVALUATIONS
# ==========================================
print("\n--- 4. BATCH IMPORTING EVALUATIONS TO SUPABASE ---")
BATCH_SIZE = 100
for i in range(0, len(evaluations_list), BATCH_SIZE):
    batch = evaluations_list[i:i + BATCH_SIZE]
    post_batch('evaluations', batch)
    print(f"Uploaded batch {i//BATCH_SIZE + 1}/{(len(evaluations_list) + BATCH_SIZE - 1)//BATCH_SIZE} ({len(batch)} records)")

print("\n🎉 ALL EXCEL DATA HAS BEEN SUCCESSFULLY IMPORTED TO SUPABASE!")
